"""Text extraction. Every extractor returns a list of PageText so page numbers
can be carried through chunking and shown as citations."""
from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

SUPPORTED_TYPES = {"pdf", "docx", "txt", "md"}


@dataclass
class PageText:
    page: int | None  # 1-based for PDFs, None for formats without pages
    text: str


class ExtractionError(Exception):
    pass


def extract_text(path: Path, file_type: str) -> list[PageText]:
    if file_type == "pdf":
        pages = _extract_pdf(path)
    elif file_type == "docx":
        pages = _extract_docx(path)
    elif file_type in {"txt", "md"}:
        pages = _extract_plain(path)
    else:
        raise ExtractionError(f"Unsupported file type: .{file_type}")

    pages = [p for p in pages if p.text.strip()]
    if not pages:
        raise ExtractionError(
            "No readable text found. Scanned PDFs need OCR before they can be indexed."
        )
    return pages


def _extract_pdf(path: Path) -> list[PageText]:
    try:
        reader = PdfReader(str(path))
        if reader.is_encrypted:
            reader.decrypt("")
        return [
            PageText(page=i, text=(page.extract_text() or ""))
            for i, page in enumerate(reader.pages, start=1)
        ]
    except Exception as exc:  # pypdf raises many different error types
        raise ExtractionError(f"Could not read PDF: {exc}") from exc


def _extract_docx(path: Path) -> list[PageText]:
    try:
        doc = DocxDocument(str(path))
    except Exception as exc:
        raise ExtractionError(f"Could not read DOCX: {exc}") from exc

    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return [PageText(page=None, text="\n".join(parts))]


def _extract_plain(path: Path) -> list[PageText]:
    try:
        return [PageText(page=None, text=path.read_text(encoding="utf-8", errors="replace"))]
    except OSError as exc:
        raise ExtractionError(f"Could not read file: {exc}") from exc
